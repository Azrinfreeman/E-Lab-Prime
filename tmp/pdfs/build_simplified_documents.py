from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCX_OUT = ROOT / "output" / "documents"
DOCX_OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
TEAL = RGBColor(17, 128, 124)
MUTED = RGBColor(92, 103, 115)
WHITE = RGBColor(255, 255, 255)


LEARNING_VALUE = [
    (
        "Pedagogy & Learning Value",
        [
            (
                "What specific learning theory is this app built upon? (e.g., Constructivism, scaffolding).",
                "The app uses constructivist and experiential learning, supported by scaffolding. Students perform an experiment, observe changes, record results, and answer a concept question. Tutorials, Professor Hakim, and real-time graphs guide them through the process.",
            ),
            (
                "How does this app measure or assess actual cognitive mastery? (Look for adaptive difficulty rather than simple rote memorization).",
                "Currently, it measures lab completion and checkpoint understanding through quizzes, XP, credits, and achievements. It does not yet use adaptive difficulty or diagnose mastery for each concept.",
            ),
            (
                "Does the app promote active learning, or is it just a glorified digital textbook?",
                "It promotes active learning. Students manipulate equipment and variables, observe simulations and graphs, record results, and answer related questions. The experience is guided rather than fully open-ended.",
            ),
            (
                "How are student errors handled?",
                "Students receive immediate feedback and can try again. The app points them back to the graph, observation, or variables instead of only saying 'incorrect.' A future improvement is feedback tailored to each wrong choice.",
            ),
        ],
    ),
    (
        "Content & Curriculum Alignment",
        [
            (
                "Is the content developed by accredited university faculty, or outsourced to third-party developers?",
                "The university team mainly developed the scientific content. Our team converted it into a game by building the code, simulations, interface, and interaction mechanics. The academic content was not outsourced to third-party content developers.",
            ),
            (
                "How does the app integrate with broader academic curriculums and learning outcomes?",
                "The app follows KSSM Form 1 Science and covers eight practical investigations. Each lab supports curriculum outcomes through variable manipulation, observation, data recording, and a concept check.",
            ),
            (
                "Is there empirical, published evidence verifying that this app improves student outcomes?",
                "There is no published evidence yet. A formal pilot using pre-tests, post-tests, and follow-up assessments is needed to measure its effect on learning outcomes.",
            ),
        ],
    ),
    (
        "Usability & Technical Quality",
        [
            (
                "Does the app feature clean, distraction-free layouts, or are users bombarded by frivolous animations and ads?",
                "The interface is focused and contains no ads. Animations, Professor Hakim, the boot sequence, and simulation effects are used for guidance and feedback rather than distraction.",
            ),
            (
                "How robust is the app's progress-tracking system?",
                "Students can see active and completed labs, achievements, XP, and credits. It tracks overall completion clearly, but it does not yet show detailed concept mastery or identify weak areas.",
            ),
        ],
    ),
    (
        "Privacy & Accessibility",
        [
            (
                "How rigorously does the app protect student data, and does it meet institution-level privacy standards?",
                "The production version will use a controlled server. Student profiles and progress will be encrypted in transit and at rest, passwords will be securely hashed, and access will be limited to authorised users. The current local storage is for prototyping, and a formal security and privacy review is required before institutional deployment.",
            ),
        ],
    ),
]


GAMIFICATION = [
    (
        "Pedagogy vs. Gamification",
        [
            (
                "How do you distinguish genuine cognitive engagement from superficial gamification, such as points or badges?",
                "Points, XP, and badges motivate students, but they do not prove learning. Real engagement happens when students change a variable, observe the effect, record evidence, and explain the result. The science task is the learning; the reward comes afterwards.",
            ),
            (
                "How does the interface manage, or accidentally increase, cognitive load for a struggling learner?",
                "The interface reduces cognitive load by showing one task at a time, with prompts, highlights, and a clear sequence. However, a 3D environment can still overwhelm some learners, so it should be tested with students and offer a simpler mode when needed.",
            ),
            (
                "How do you account for differences in device access, screen size, and home connectivity?",
                "The research should record device and internet conditions, test lower-end devices, and provide school devices or scheduled school sessions. An offline-first design allows the lab to run without internet and synchronise progress later.",
            ),
            (
                "What is your ethical framework for analysing student data without exploiting privacy?",
                "We collect only data needed for learning and research. Because the learners are minors, we require appropriate consent, pseudonymisation, encryption, limited access, retention rules, and a deletion process. Student data must never be used for commercial profiling.",
            ),
            (
                "How does your model account for the teacher's pedagogical role?",
                "The app supports the teacher; it does not replace them. Teachers introduce the goal, ask for predictions, guide discussion, and correct misconceptions. Research should compare teacher-supported use with app-only use.",
            ),
            (
                "How will you prove that learning gains last for 12 months rather than reflecting a novelty effect?",
                "Use a baseline test, an immediate post-test, and follow-ups at three, six, and twelve months, together with a comparison group. Sustained understanding and knowledge transfer, not initial enjoyment, are the real evidence.",
            ),
            (
                "If adaptive learning or AI feedback is used, how will you detect algorithmic bias?",
                "The current app does not use AI or adaptive algorithms. If they are added later, outcomes should be audited across student groups and devices, with human oversight and a teacher override.",
            ),
            (
                "How do you manage the tension between commercial retention and independent mastery?",
                "Success is measured by learning, transfer, and confidence, not screen time. The app has no ads or commercial retention loops. The goal is for students to understand the science independently.",
            ),
            (
                "How does your methodology measure deep learning instead of rote multiple-choice performance?",
                "Deep learning is measured through graph interpretation, prediction, explanation, accurate data recording, and transfer to a new situation. Multiple-choice questions can support checking, but cannot prove deep understanding on their own.",
            ),
            (
                "What developmental psychology framework justifies the required screen time, attention, and motor skills?",
                "The design targets Form 1 learners and draws on Piaget's formal operational stage, Vygotsky's scaffolding, and Cognitive Load Theory. Controls are simple, and screen time should be teacher-led, divided into short segments, and include breaks.",
            ),
            (
                "What is the main issue with using this app?",
                "The main challenge is moving from a polished prototype to full deployment. The next steps are secure server accounts, progress synchronisation, stronger mastery analytics, lower-end device testing, and formal learning evaluation.",
            ),
            (
                "Is the app online or offline?",
                "The prototype works offline. The production version should be hybrid: labs run offline, while accounts and progress synchronise securely when internet is available.",
            ),
            (
                "Is there a potential for the app to crash?",
                "Any 3D app can crash, especially on lower-end devices or during network failures. The current Unity Editor has no active errors, but device testing, quality settings, safe saves, crash logs, and recovery features are still required.",
            ),
            (
                "Who is responsible for crashes or other issues?",
                "The development team is responsible for code, bugs, server reliability, security, releases, and support. The university content team is responsible for scientific accuracy and KSSM alignment. Teachers guide classroom use and report issues; students are not responsible for technical failures.",
            ),
        ],
    ),
]


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_numbering_definition(document):
    numbering = document.part.numbering_part.element
    existing_abstract = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abstract, default=-1) + 1
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    level.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = BLUE
    heading1.paragraph_format.space_before = Pt(18)
    heading1.paragraph_format.space_after = Pt(10)
    heading1.paragraph_format.keep_with_next = True

    question = doc.styles.add_style("QA Question", WD_STYLE_TYPE.PARAGRAPH)
    question.base_style = normal
    question.font.name = "Calibri"
    question._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    question._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    question.font.size = Pt(11.5)
    question.font.bold = True
    question.font.color.rgb = NAVY
    question.paragraph_format.space_before = Pt(9)
    question.paragraph_format.space_after = Pt(3)
    question.paragraph_format.line_spacing = 1.15
    question.paragraph_format.keep_with_next = True

    answer = doc.styles.add_style("QA Answer", WD_STYLE_TYPE.PARAGRAPH)
    answer.base_style = normal
    answer.font.name = "Calibri"
    answer._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    answer._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    answer.font.size = Pt(11)
    answer.font.color.rgb = RGBColor(35, 44, 54)
    answer.paragraph_format.left_indent = Inches(0.375)
    answer.paragraph_format.space_after = Pt(8)
    answer.paragraph_format.line_spacing = 1.25
    answer.paragraph_format.keep_together = True


def build_document(title, filename, sections):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    num_id = add_numbering_definition(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("E-LAB PRIME  |  PRESENTATION Q&A")
    set_run_font(hr, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.keep_with_next = True
    title_r = title_p.add_run(title)
    set_run_font(title_r, size=24, bold=True, color=NAVY)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(16)
    subtitle_r = subtitle_p.add_run("Simplified presentation answers")
    set_run_font(subtitle_r, size=10.5, italic=True, color=MUTED)

    for section_title, items in sections:
        heading = doc.add_paragraph(section_title, style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        for question_text, answer_text in items:
            question = doc.add_paragraph(style="QA Question")
            apply_numbering(question, num_id)
            question.add_run(question_text)

            answer = doc.add_paragraph(style="QA Answer")
            label = answer.add_run("Answer: ")
            set_run_font(label, size=10.5, bold=True, color=TEAL)
            response = answer.add_run(answer_text)
            set_run_font(response, size=11, color=RGBColor(35, 44, 54))

    output = DOCX_OUT / filename
    doc.save(output)
    return output


def main():
    outputs = [
        build_document(
            "Pedagogy & Learning Value",
            "Pedagogy and Learning Value - Simplified.docx",
            LEARNING_VALUE,
        ),
        build_document(
            "Pedagogy vs. Gamification",
            "Pedagogy vs Gamification - Simplified.docx",
            GAMIFICATION,
        ),
    ]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
